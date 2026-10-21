"""중간 평가 기획안 ④ ⑤ Word 문서 두 개 (형식: 참고 문서 '도서 대여 시스템 만들기.docx' / 같은 내용의 Google 문서)
- ④ 노포 천년 미식가 만들기.docx : 그림이 많은 전체 개발 기록(기획 → 디자인 → 데이터 → 개발 → 테스트·버그 → 세부 수정 → 배포)
- ⑤ 노포 천년 미식가 개발 기록 (Google 문서용).docx : 같은 구성을 번호 장·표로 간추린 문서 (Google 문서에서 열어 쓰도록)
참고 docx를 틀(스타일·번호·글꼴)로 열어 본문만 바꾼다. 실행: python tools/mid_eval/build_docx.py"""
import copy
import glob
import io
import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(__file__))
import assets  # noqa: E402

REF = glob.glob(os.path.join(assets.ROOT, 'ex_copy', '*.docx'))[0]
OUT1 = os.path.join(assets.OUT_DIR, '노포 천년 미식가 만들기.docx')
OUT2 = os.path.join(assets.OUT_DIR, '노포 천년 미식가 개발 기록 (Google 문서용).docx')
FONT = 'Arial Unicode MS'
DOWNLOAD = 'https://github.com/ohdangchan/NopoCheonnyeon/releases/latest/download/nopo-cheonnyeon.apk'
REPO = 'https://github.com/ohdangchan/NopoCheonnyeon'


class Doc:
    def __init__(self):
        self.d = Document(REF)
        body = self.d.element.body
        for el in list(body):
            if el.tag != qn('w:sectPr'):
                body.remove(el)
        # 참고 문서의 그림은 모두 버린다(용량)
        for rid, rel in list(self.d.part.rels.items()):
            if rel.reltype.endswith('/image'):
                del self.d.part.rels[rid]
        # 본문 폭: A4, 여백 1440 → 약 15.9cm
        self.width_cm = 15.9

    # ---- 글자
    def _run(self, p, text, bold=False, size=None, color=None, mono=False, underline=False):
        r = p.add_run(text)
        rpr = r._r.get_or_add_rPr()
        f = OxmlElement('w:rFonts')
        name = 'Consolas' if mono else FONT
        for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
            f.set(qn(a), name if a != 'w:eastAsia' or not mono else FONT)
        rpr.insert(0, f)
        if bold:
            r.bold = True
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if underline:
            r.underline = True
        return r

    def para(self, text='', style=None, bold=False, size=None, color=None, mono=False, space_after=None):
        p = self.d.add_paragraph(style=style)
        if text:
            self._run(p, text, bold=bold, size=size, color=color, mono=mono)
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        return p

    def rich(self, parts, style=None):
        """parts: [(text, {bold, color, mono})]"""
        p = self.d.add_paragraph(style=style)
        for t, o in parts:
            self._run(p, t, **o)
        return p

    def title(self, text):
        p = self.d.add_paragraph(style='Title')
        self._run(p, text, size=26)
        return p

    def h1(self, text):
        p = self.d.add_paragraph(style='Heading1')
        self._run(p, text, size=20)
        return p

    def h2(self, text):
        p = self.d.add_paragraph(style='Heading2')
        self._run(p, text, size=16)
        return p

    def h3(self, text):
        p = self.d.add_paragraph(style='Heading3')
        self._run(p, text, size=14)
        return p

    def link(self, label, url):
        p = self.d.add_paragraph()
        self._run(p, label + ' ')
        self._run(p, url, color='1155CC', underline=True)
        return p

    def bullet(self, text, level=0, bold_prefix=None):
        p = self.d.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        num = OxmlElement('w:numPr')
        il = OxmlElement('w:ilvl'); il.set(qn('w:val'), str(level))
        ni = OxmlElement('w:numId'); ni.set(qn('w:val'), '2')
        num.append(il); num.append(ni)
        ppr.append(num)
        ind = OxmlElement('w:ind'); ind.set(qn('w:left'), str(720 + 360 * level)); ind.set(qn('w:hanging'), '360')
        ppr.append(ind)
        if bold_prefix:
            self._run(p, bold_prefix, bold=True)
        self._run(p, text)
        return p

    def code(self, lines):
        """회색 상자 안의 코드·프롬프트 (1칸 표)"""
        t = self.d.add_table(rows=1, cols=1)
        self._borders(t, 'CCCCCC')
        c = t.rows[0].cells[0]
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'F3F5F3')
        c._tc.get_or_add_tcPr().append(shd)
        c.paragraphs[0].text = ''
        for i, ln in enumerate(lines):
            p = c.paragraphs[0] if i == 0 else c.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            self._run(p, ln, mono=True, size=9.5)
        self.para('')

    # ---- 표
    def _borders(self, t, color='999999'):
        tblPr = t._tbl.tblPr
        b = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement('w:' + side)
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '8'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
            b.append(e)
        tblPr.append(b)

    def table(self, rows, widths=None, header=True, size=10):
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        self._borders(t)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.rows[i].cells[j]
                if widths:
                    c.width = Cm(widths[j])
                c.paragraphs[0].text = ''
                self._run(c.paragraphs[0], str(val), bold=(header and i == 0), size=size)
                if header and i == 0:
                    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'E8EEEC')
                    c._tc.get_or_add_tcPr().append(shd)
        self.para('')
        return t

    def images(self, items, per_row=3, caption=True):
        """items: [(PIL 이미지, 설명)] — 표 안에 그림을 나란히 (원본 문서의 화면 모음 표와 같은 모양)"""
        rows = (len(items) + per_row - 1) // per_row
        t = self.d.add_table(rows=rows * (2 if caption else 1), cols=per_row)
        self._borders(t, 'DDDDDD')
        w = (self.width_cm - 0.6) / per_row
        for k, (im, cap) in enumerate(items):
            r, c = divmod(k, per_row)
            cell = t.rows[r * (2 if caption else 1)].cells[c]
            buf = io.BytesIO(assets.png_bytes(im))
            cell.paragraphs[0].alignment = 1
            cell.paragraphs[0].add_run().add_picture(buf, width=Cm(w - 0.3))
            if caption:
                cc = t.rows[r * 2 + 1].cells[c]
                cc.paragraphs[0].alignment = 1
                cc.paragraphs[0].text = ''
                self._run(cc.paragraphs[0], cap, size=9)
        self.para('')

    def image(self, im, width_cm=6):
        buf = io.BytesIO(assets.png_bytes(im))
        p = self.d.add_paragraph()
        p.add_run().add_picture(buf, width=Cm(width_cm))
        return p

    def save(self, path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.d.save(path)
        print('saved', path, os.path.getsize(path))


S = {k: v[1] for k, v in assets.SCREENS.items()}
D = {k: v[1] for k, v in assets.DEVICE.items()}

REQUESTS = [
    ('1', '지도에서 가게를 볼 때 사진이나 상호를 눌러도 가게 상세로 가게', '하단 시트의 사진·이름 영역을 눌러 가게 상세로 이동 (자세히 보기 버튼과 같은 동작)'),
    ('2', '다른 가게를 보다가 "내 위치"를 누르면 내 위치로 이동, 위치 잡는 속도 개선', '내 위치 버튼이 지도를 내 위치로 재중심, 마지막 위치를 먼저 쓰는 빠른 측정(quickFix) 추가'),
    ('3', '화면별 미리보기와 Firestore 호출 시험', '화면별 @Preview 추가, Firestore 호출 단위 시험 작성'),
    ('4', '도장 애니메이션, 칭호 100년 단위, 공유 카드에 QR', '도장 연출(380ms·진동), 칭호 11단계, 패스포트 공유 카드에 다운로드 QR(ZXing)'),
    ('5', '리뷰에 갤러리 사진 최대 3장(파일 저장 없이)', '카메라·시스템 사진 선택기로 3장, 720px WebP로 줄여 문자열로만 저장 (CLAUDE.md 예외 규칙 추가)'),
    ('6', '시연용 사진이 너무 구림 → 실제 음식 사진', '무료 사진 사이트에서 메뉴별 사진을 받아 앱에 포함하고 출처 목록(data/demo_photo_credits.json) 기록'),
    ('7', '칭호 배지 색이 바뀌는지 알 수 없다 / 칭호 연출이 더 화려했으면', '칭호 도감 화면 추가, 칭호 승급 연출(광선·색종이·폭죽) 강화'),
    ('8', '내 리뷰를 쓰고 또 쓰면 전 리뷰가 사라진다', '리뷰 문서를 가게_사용자_슬롯(1~5)으로 바꿔 한 가게에 여러 리뷰, 서버 규칙·시험·PRD 수정'),
    ('9', '도장을 쾅 찍고 난 후 애니메이션이 뜨게, 애니메이션을 동영상으로', '도장 → 충격 고리 → 폭죽 → 승급 순서로 정리, 화면 녹화 docs/media/stamp_ceremony_demo.mp4'),
    ('10', '새 폰에 앱을 설치해서 계속 개발', 'Galaxy S9+(Android 10)에 설치·확인, 구형 기기 문제 점검'),
    ('11', 'AI 메뉴 추천 룰렛(내 리뷰·메뉴·음식 사진 → 메뉴 3개, 최대 30개)', 'MenuRoulette 규칙, Gemini 호출, 처음 사용 전 동의, 룰렛 화면과 사진 릴'),
    ('12', '가게별 리뷰를 AI가 분석해 특장점 제공', '한 줄 요약·특장점·추천 메뉴·아쉬운 점 카드, 같은 리뷰는 캐시'),
    ('13', '재방문율을 비교해 단골 충성도 100점 만점', '기준 점수는 앱이 계산하고 AI 보정은 ±12점으로 제한'),
    ('14', '가게 소개·리뷰 AI 번역', '번역 막대 추가(영어 등), 개수·순서 검증'),
    ('15', '지도 위 AI 챗봇(말로 찾기·천년 경로·오늘의 코스), 대화 기록 유지', 'AI 노포 도우미 화면, 후보 가게 압축, 기록은 기기에 저장'),
    ('16', '영수증 스캔 레이저·형광펜 효과, AI 실패 팝업, 테스트 영수증', '스캔 연출(레이저 → 글자별 형광펜), AI 촬영 도우미 팝업, 시험 영수증 생성 스크립트와 시나리오 문서'),
    ('17', '간판·메뉴판 인증에도 같은 효과, 리뷰 남기기는 방문 인증 후', '간판 인증 스캔 연출과 원인·대안 안내, 리뷰 버튼 → 인증 화면 → 편집기 흐름'),
    ('18', '검색창과 AI 도우미 합치기, 예시 질문 보이기', '검색창 한 개 + "AI" 버튼, 문장형 입력 시 도우미 제안, 예시 질문 카드'),
    ('19', '서울 가게마다 리뷰 10개 더, 내 사진 리뷰 10개, 룰렛에 내 사진 슬라이딩', '시연 리뷰 4,956개, 내 사진 리뷰 10개, 사진 릴 효과'),
    ('20', '시나리오가 뒤죽박죽 → 전부 새로, 시험용 가게 복원', 'demo-guide.md·demo_scenarios.html 정리, 테스트 가게 재생성'),
    ('21', '최종 릴리스 APK + QR 다운로드, 저장소 공개', 'GitHub 릴리스 v0.1.0-demo, QR, 저장소 공개 전환'),
    ('22', 'QA 검수 문서', 'docs/qa-final-2026-10-21.md (291 테스트·기능표·보안·결함·체크리스트)'),
    ('23', 'APK에서 API 키 제거', 'Cloud Functions 안을 버리고 Firebase AI Logic으로 전환, App Check 적용 해제, 25초 제한·예비 모델, 키 없는 릴리스로 교체'),
    ('24', '앱 아이콘 제작', '청록 간판 + 붉은 도장 적응형 아이콘, 생성 스크립트 tools/gen_app_icon.py'),
]


def full():
    d = Doc()
    d.title('노포 천년 미식가 만들기')
    d.para('')
    d.para('기획부터 디자인, 개발, 실기기 검증, 배포까지 Claude Code와 Google 서비스로 진행')
    d.bullet(' : UI 디자인 (시안 51화면)', bold_prefix='Stitch ')
    d.bullet(': 노포 데이터 수집·변환 (Python 스크립트)', bold_prefix='Python 수집 스크립트')
    d.bullet(': 개발 (Android Studio와 함께)', bold_prefix='Claude Code')
    d.bullet(': 로그인 · 데이터베이스 · AI(Gemini) · 크래시 수집', bold_prefix='Firebase')
    d.bullet(': 릴리스 APK 배포 · QR 다운로드', bold_prefix='GitHub')
    d.link('완성된 앱 내려받기(APK):', DOWNLOAD)
    d.link('소스 저장소:', REPO)
    d.image(assets.qr(300), 4)
    d.images([(D['dev_map'](380), '지도 탐색'), (D['dev_passport'](380), '내 패스포트'), (D['dev_chat'](380), 'AI 노포 도우미')])

    # ---------------- 기획
    d.h1('기획')
    d.para('만들고싶은 시스템에 대해서 정리')
    d.para('# 노포 천년 미식가', bold=True)
    d.para('Android 앱 (Kotlin, Jetpack Compose) · 가 본 노포의 업력(년)을 세월로 모아 1,000년을 채우는 미식 패스포트')
    d.para('## 핵심 규칙', bold=True)
    d.bullet('방문 인증 1회 = 그 가게의 업력(올해 − 창업연도)만큼 세월 적립. 창업연도를 모르면 30년')
    d.bullet('같은 가게는 처음 한 번만 적립, 다시 가면 단골 횟수만 증가')
    d.bullet('세월 100년마다 칭호 한 단계 (11단계: 첫술 식객 … 천년 미식가)')
    d.bullet('도장 등급: 업력 30~49 일반(청록), 50~99 붉은, 100 이상 금빛')
    d.para('## 필수요소', bold=True)
    d.bullet('계정: 누가 인증·리뷰했는지 알아야 하므로 회원가입·로그인·로그아웃이 필요. 이메일·비밀번호·닉네임(2~12자, 중복 불가)')
    d.bullet('방문 인증: 영수증이 기본(위치 정보 사용 안 함). 안 되면 가게 30m 이내에서 간판·메뉴판 사진')
    d.bullet('리뷰: 그 가게를 방문 인증한 사람만 별점·글·사진 작성')
    d.bullet('로그인 게이트: 로그인 전에는 온보딩·로그인·가입 화면만')
    d.para('## 화면', bold=True)
    d.para('### 지도 탐색 화면', bold=True)
    d.bullet('네이버 지도 + 마커(미방문 청록, 방문 붉은 도장, 업력 100년 이상 금빛 테두리)')
    d.bullet('필터 칩: 음식점만 / 전체, 백년가게 / 오래가게, 안 가 본 곳. 검색: 가게명·지역·메뉴')
    d.bullet('마커를 누르면 하단 시트 → 사진·이름·자세히 보기 버튼 모두 가게 상세로 이동')
    d.para('### 가게 상세 화면', bold=True)
    d.bullet('업력 큰 숫자, 가게 이야기(소개글), 영업시간, 대표 메뉴, 평균 평점, 리뷰 목록(20개씩)')
    d.bullet('"인증하면 세월 +N년" 보상을 먼저 노출, 방문 인증하기 버튼')
    d.para('### 방문 인증 화면', bold=True)
    d.bullet('영수증 인증: 앱 안 카메라로 촬영 → 기기 안 글자 인식 → 전화번호 일치 또는 상호 유사도 0.70 이상, 결제일시 24시간 이내')
    d.bullet('가게 사진 인증: 정밀 위치 15초 이내 측정 → 거리 30m·정확도 30m 이하 → 간판·메뉴판 글자 확인')
    d.bullet('결과: 도장 연출(380ms·진동), 세월 +N년, 칭호 승급')
    d.para('### 패스포트·리뷰·AI 화면', bold=True)
    d.bullet('패스포트: 세월 게이지, 도장 격자, 도장 상세, 내 리뷰, 공유 카드(QR)')
    d.bullet('AI 노포 도우미: 말로 찾기 · 천년 경로 · 오늘의 코스, 메뉴 룰렛, 후기 분석, 단골 충성도')

    # ---------------- 디자인
    d.h1('디자인')
    d.h2('Stitch에서 디자인 만들기')
    d.para('기획한 내용을 프롬프트에 입력해 51개 화면을 생성하고 기능 패키지별로 정리 (design/stitch/screens)')
    d.images([(S['login'](380), '로그인'), (S['map'](380), '지도 탐색'), (S['store'](380), '가게 상세')])
    d.h2('기준 문서와 차이 정리')
    d.para('시안이 PRD와 다른 부분은 design/stitch/PRD-차이점.md에 화면별로 적고, 화면을 만들기 전에 반드시 읽게 했다. 기준 우선순위: PRD → 아키텍처 → 스토리·체크리스트 → 기획안 → Stitch 시안.')
    d.h2('Stitch에서 세부 조정 후')
    d.para('색·글꼴은 시안 HTML 값이 아니라 테마 토큰(sign, primary, stamp, brass, paper)으로 옮김. 다크 모드 한 벌 추가.')
    d.images([(S['receipt'](380), '영수증 인증'), (S['photo'](380), '가게 사진 인증'), (S['stamp'](380), '도장 연출'),
              (S['passport'](380), '내 패스포트'), (S['review'](380), '리뷰 작성'), (S['title'](380), '칭호 승급')])

    # ---------------- 데이터 수집
    d.h1('데이터 수집')
    d.h2('공식 자료에서 필수 데이터 수집')
    d.para('앱을 구성하는데 필요한 노포 데이터 모으기')
    d.para('수집 순서:')
    d.code(['data/raw/*.jsonl                수집 원본 (소상공인마당 백년가게, 서울시 오래가게)',
            'tools/import_official_list.py   공식 지정 리스트 CSV와 대조해 빠진 가게 보충',
            'tools/build_catalog.py          → data/nopo_stores.csv, app/src/main/assets/stores.json',
            'tools/geocode.py                주소 → 좌표 (네이버/카카오 이름 검색 보강)',
            'tools/fetch_stories.py          가게 이야기(소개글) 수집'])
    d.h2('결과 확인')
    d.table([
        ['항목', '수', '설명'],
        ['전체 카탈로그', '1,277곳', '백년가게 1,138 + 서울 오래가게 151 (겹치는 12곳 포함)'],
        ['음식점업', '834곳', '지도 기본 필터'],
        ['좌표 확보', '1,181곳', '좌표 없는 가게는 목록·검색에서 "위치 준비 중"'],
        ['가게 이야기', '1,138곳', '소상공인마당 소개글 (출처 표시)'],
        ['평균 업력', '40.0년', '창업연도를 아는 1,237곳 기준, 100년 이상 8곳'],
        ['창업연도 미상', '40곳', '업력 30년(공식 지정 최소 기준)으로 계산, "업력 확인 중" 표시'],
    ], widths=[3.6, 2.4, 9.9])

    # ---------------- 개발
    d.h1('개발')
    d.h2('개발 폴더 준비')
    d.para('폴더 이름: NopoCheonnyeon (GitHub: ohdangchan/NopoCheonnyeon)')
    d.code(['NopoCheonnyeon/', '├─ CLAUDE.md                 Claude Code 작업 지침 (반드시 지킬 것)',
            '├─ docs/                     PRD·아키텍처·스토리·체크리스트·기획안·QA·인수인계',
            '├─ design/stitch/            Stitch 시안 51화면 + PRD 차이점',
            '├─ data/ · tools/            수집 데이터와 변환·생성 스크립트',
            '├─ firebase/                 firestore.rules, indexes',
            '└─ app/src/main/java/com/nopo/cheonnyeon/',
            '     feature/  (화면+ViewModel)  domain/ (규칙)  data/ (Room·Firestore·AI)  platform/ (카메라·위치·OCR)'])
    d.h2('Claude Code에게 준 규칙 (CLAUDE.md 요약)')
    d.code(['- 사용자에게 보이는 문구는 strings.xml, 색은 테마 토큰만 사용',
            '- 인증 기준값은 VerificationConfig 한 곳에만',
            '- 영수증 인증은 위치를 쓰지 않는다. 영수증 이미지는 저장하지 않는다',
            '- 위치 좌표와 영수증 내용은 서버에 저장하지 않고 로그에도 남기지 않는다',
            '- 키·비밀번호를 코드나 로그에 쓰지 않는다',
            '- 기기 기능·서버 호출은 withTimeout과 AppError 변환으로 감싼다',
            '- 아키텍처 문서에 없는 라이브러리는 이유를 말하고 먼저 묻는다',
            '- 한 번에 스토리 하나: 계획 5줄 → 테스트 먼저 → 빌드·설치 → nopo-qa 검증'])
    d.h2('기술 스택')
    d.table([
        ['영역', '선택'],
        ['언어·UI', 'Kotlin 2.2, Jetpack Compose, Material 3, MVVM + 단방향 흐름, 수동 주입(AppContainer)'],
        ['지도·위치·카메라', '네이버 지도 SDK 3.24.0, Fused Location, CameraX'],
        ['글자 인식', 'ML Kit 한국어 텍스트 인식 (앱에 포함, 오프라인)'],
        ['서버', 'Firebase Authentication, Cloud Firestore(보안 규칙), Crashlytics'],
        ['기기 저장', 'Room, DataStore, WorkManager(대기열 전송)'],
        ['AI', 'Gemini (Firebase AI Logic), 25초 제한 + 예비 모델'],
        ['그 밖에', 'Coil, kotlinx.serialization, ZXing(QR 생성)'],
    ], widths=[3.6, 12.3])
    d.h2('Firebase 연동')
    d.bullet('Authentication: 이메일/비밀번호 로그인 사용 설정')
    d.bullet('Cloud Firestore: 서울 리전, 보안 규칙 배포 (cd firebase && firebase deploy --only firestore:rules)')
    d.bullet('규칙 요점: 방문 문서가 있을 때만 리뷰 쓰기, 필드 화이트리스트, 사진 크기, 닉네임 일치, 서버 시각')

    # ---------------- 테스트·버그
    d.h1('테스트 및 버그 수정 (실사용)')
    d.h2('자동 테스트')
    d.bullet('단위 테스트 291개 통과, 실패 0 (testDebugUnitTest)')
    d.bullet('린트 오류 0, 경고 25 (lintDebug)')
    d.bullet('서버 규칙 에뮬레이터 시험 18개는 작성했으나 이 PC에서 실행 못함(환경 문제)')
    d.h2('실기기 시험 (Galaxy S9+ Android 10, Z Fold Android 17)')
    d.images([(D['dev_map'](380), '릴리스 앱: 지도'), (D['dev_passport'](380), '릴리스 앱: 패스포트'), (D['dev_chat'](380), '릴리스 앱: AI 도우미')])
    d.h2('문제와 해결')
    d.table([
        ['증상', '원인', '해결'],
        ['릴리스 앱이 시작하자마자 종료', 'R8 최적화가 Crashlytics와 충돌 (IllegalAccessError)', '최적화는 끄고 축소·난독화만 사용 (-dontoptimize)'],
        ['내 위치가 늦게 잡힘', '정확한 위치를 새로 받을 때까지 대기', '마지막 위치를 먼저 쓰는 quickFix + 내 위치 버튼 재중심'],
        ['리뷰를 또 쓰면 이전 리뷰가 사라짐', '문서 ID가 가게_사용자 하나', '슬롯 1~5로 변경, 서버 규칙·시험 수정'],
        ['Gemini 키가 공개 APK에 포함', '앱이 키를 직접 가짐', 'Firebase AI Logic으로 이전, 키 제거, 기존 키 폐기'],
        ['"App Check token is invalid"', 'AI Logic에 App Check 적용이 켜져 있음', '콘솔 App Check → APIs에서 적용 해제 (반영 약 10분)'],
        ['AI 응답이 40초 넘게 멈춤', '기본 모델이 붐빔(high demand)', '요청 25초 제한 + 예비 모델 한 번 재시도'],
    ], widths=[4.2, 5.4, 6.3])
    d.h2('미확인으로 남긴 항목')
    d.para('영수증·간판 스캔 연출의 실기기 화면, AI 촬영 도우미 팝업 화면, 강화된 서버 규칙으로의 실제 리뷰 저장, 리뷰 사진 카메라·앨범 업로드는 사용자가 직접 확인해야 하는 항목으로 QA 문서(docs/qa-final-2026-10-21.md)에 "미확인(사용자)"로 분리했다.')

    # ---------------- 세부 수정
    d.h1('세부 수정 내역')
    d.table([['번호', '요청', '결과']] + [list(r) for r in REQUESTS], widths=[1.2, 6.4, 8.3], size=9.5)

    # ---------------- 배포
    d.h1('배포')
    d.h2('릴리스 APK')
    d.bullet('assembleRelease: R8 축소·난독화, 릴리스 키 서명, arm64만 포함 (약 88MB)')
    d.bullet('폰에 덮어 설치: adb install -r app/build/outputs/apk/release/app-release.apk')
    d.bullet('GitHub 공개 릴리스 v0.1.0-demo에 nopo-cheonnyeon.apk 업로드, 올린 파일과 빌드 결과의 크기가 같은지 확인')
    d.link('내려받기:', DOWNLOAD)
    d.image(assets.qr(300), 4)
    d.h2('배포 정보')
    d.table([
        ['항목', '내용'],
        ['앱 이름 / 패키지', '노포 천년 미식가 / com.nopo.cheonnyeon (디버그와 릴리스 동일)'],
        ['Firebase 프로젝트 ID', 'nopocheonnyeon (Authentication, Firestore 서울 리전, Crashlytics, AI Logic)'],
        ['저장소', REPO + ' (공개)'],
        ['최신 릴리스', 'v0.1.0-demo (arm64, 약 88MB)'],
        ['개발 기간', '2026-10-19 ~ 2026-10-21 (커밋 106개)'],
    ], widths=[4.4, 11.5])
    d.h2('아이콘')
    d.image(assets.logo(400), 4)
    d.para('청록 간판 바탕 + 붉은 도장("노포"). python tools/gen_app_icon.py 로 모든 해상도를 생성.')
    d.save(OUT1)


def concise():
    d = Doc()
    d.title('노포 천년 미식가 개발 기록')
    d.para('')
    d.h1('문서 개요')
    d.para('이 문서는 안드로이드 앱 "노포 천년 미식가"를 기획·디자인·개발·검증·배포한 과정을 정리한 작업 로그입니다. 그림이 많은 전체 기록은 "노포 천년 미식가 만들기.docx"에 있고, 이 문서는 Google 문서에 올려 공유하기 쉽게 장·표로 간추렸습니다.')
    d.h2('1. 프로젝트 개요')
    d.bullet('가 본 노포의 업력(년)을 세월로 모아 1,000년을 채우는 안드로이드 앱', bold_prefix='목표: ')
    d.bullet('Stitch(UI 시안), Python 스크립트(데이터), Claude Code(개발), Firebase(로그인·DB·AI), GitHub(배포)', bold_prefix='사용 도구: ')
    d.bullet('Kotlin, Jetpack Compose, 네이버 지도 SDK, ML Kit, Room, Firestore, Gemini(Firebase AI Logic)', bold_prefix='핵심 기술: ')
    d.link('앱 내려받기:', DOWNLOAD)
    d.link('저장소:', REPO)

    d.h2('2. 기획')
    d.bullet('1회 인증 = 가게 업력만큼 세월, 100년마다 칭호 11단계, 도장 3등급', bold_prefix='세월 규칙: ')
    d.bullet('영수증(위치 미사용) → 안 되면 가게 30m 안에서 간판·메뉴판 사진', bold_prefix='인증: ')
    d.bullet('그 가게를 방문 인증한 사용자만 별점·글·사진(3장) 리뷰', bold_prefix='리뷰: ')
    d.bullet('이메일 가입·로그인·닉네임, 로그인 전에는 가입 화면만', bold_prefix='계정: ')
    d.bullet('지도·목록, 가게 상세, 인증, 도장 연출, 패스포트, 리뷰, 프로필, AI 노포 도우미', bold_prefix='화면: ')

    d.h2('3. 디자인 (Stitch)')
    d.bullet('기획 내용을 프롬프트로 입력해 51개 화면 생성, 기능 패키지별로 정리')
    d.bullet('"골목 간판과 인주": 간판 청록 #17504A, 포인트 청록 #1F6B63, 인주 빨강 #C23B2A, 놋쇠 #D9A22C, 종이 #EDF0EA')
    d.bullet('시안이 PRD와 다른 부분은 PRD-차이점.md로 정리, 색은 값이 아니라 토큰으로 옮김')
    d.images([(S['map'](340), '지도'), (S['store'](340), '가게 상세'), (S['passport'](340), '패스포트')])

    d.h2('4. 데이터 수집')
    d.bullet('소상공인마당 백년가게 + 서울시 오래가게 → 1,277곳 (음식점 834, 좌표 1,181, 이야기 1,138)')
    d.bullet('파이프라인: data/raw → tools/build_catalog.py → stores.json (앱 assets)')
    d.bullet('창업연도 미상 40곳은 업력 30년(최소 기준)으로 계산하고 "업력 확인 중" 표시')

    d.h2('5. 개발 (Claude Code)')
    d.h3('5-1. 1차 개발')
    d.bullet('1일차: 빈 앱에 지도·카메라·글자 인식·위치·Firebase 진단 화면으로 실기기 검증 → 최종 앱의 진단 화면으로 유지')
    d.bullet('로그인 게이트, 지도, 가게 상세, 영수증·사진 인증, 도장·세월·칭호, 패스포트, 리뷰·프로필 구현')
    d.h3('5-2. 추가 규칙 반영')
    d.bullet('같은 영수증 재사용 금지(해시), 하루 5곳·같은 가게 1회, 오프라인 대기열 → 연결 시 자동 전송')
    d.bullet('리뷰 여러 개(슬롯 1~5), 리뷰 남기기 → 방문 인증 → 편집기, 위치 확인 도장')
    d.h3('5-3. AI 기능 (Gemini)')
    d.bullet('AI 노포 도우미(말로 찾기·천년 경로·오늘의 코스), 메뉴 룰렛 30개, 후기 분석, 단골 충성도, 번역, 촬영 도우미')
    d.bullet('후보 가게 안에서만 추천, 숫자는 앱이 계산, 처음 사용 전 전송 동의')
    d.h3('5-4. Firebase 연동')
    d.bullet('Authentication·Firestore(서울)·Crashlytics, 보안 규칙(필드 화이트리스트·방문 문서 필요)')
    d.bullet('키 없이 Gemini 호출: Firebase AI Logic (App Check 적용은 시연용으로 해제, 2026-11-02부터 의무)')

    d.h2('6. 테스트 및 버그 수정 (실사용)')
    d.h3('6-1. 자동 테스트')
    d.bullet('단위 테스트 291개 통과, 린트 오류 0')
    d.h3('6-2. 실기기 (Galaxy S9+, Z Fold)')
    d.bullet('릴리스 서명 APK로 지도·패스포트·AI 도우미 확인')
    d.h3('6-3. 주요 버그')
    d.table([
        ['증상', '원인', '해결'],
        ['릴리스 앱 시작 시 종료', 'R8 최적화 ↔ Crashlytics', '-dontoptimize'],
        ['위치가 늦게 잡힘', '새 위치 측정 대기', 'quickFix(마지막 위치 우선)'],
        ['리뷰 또 쓰면 이전 리뷰 사라짐', '문서 ID 하나', '슬롯 1~5'],
        ['공개 APK에 AI 키', '앱이 키를 직접 가짐', 'Firebase AI Logic으로 이전'],
        ['AI 응답 40초 지연', '모델 붐빔', '25초 제한 + 예비 모델'],
    ], widths=[4.8, 5.2, 5.9])
    d.h3('6-4. 미확인')
    d.bullet('영수증·간판 스캔 연출 화면, AI 촬영 도우미 팝업, 강화된 규칙으로의 실제 리뷰 저장, 사진 리뷰 업로드(사용자 확인 필요), 서버 규칙 에뮬레이터 시험(환경 문제)')

    d.h2('7. 세부 수정 내역')
    d.table([['번호', '요청', '결과']] + [list(r) for r in REQUESTS], widths=[1.2, 6.4, 8.3], size=9.5)

    d.h2('8. 배포 정보')
    d.bullet('GitHub 공개 릴리스 v0.1.0-demo (arm64, 약 88MB), QR 다운로드', bold_prefix='배포 방식: ')
    d.bullet('com.nopo.cheonnyeon (디버그·릴리스 동일), 릴리스 키 서명', bold_prefix='패키지: ')
    d.bullet('Firebase 프로젝트 nopocheonnyeon — Authentication, Firestore(서울), Crashlytics, AI Logic', bold_prefix='백엔드: ')
    d.bullet('2026-10-19 ~ 2026-10-21, 커밋 106개', bold_prefix='개발 기간: ')
    d.image(assets.qr(300), 4)
    d.save(OUT2)


if __name__ == '__main__':
    full()
    concise()
