# -*- coding: utf-8 -*-
"""content.py 의 블록을 .docx 와 .html 로 그린다.
실행: python -X utf8 tools/guide_v2/build.py
출력: docs/v2/report/노포패스포트_버전2_만들기.docx, .html
필요: pip install python-docx pillow
"""
import base64
import html as H
import io
import os
import subprocess
import sys

from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE)
import content  # noqa: E402

SCREENS = os.path.join(ROOT, 'docs', 'v2', 'screens')
IMG = os.path.join(ROOT, 'docs', 'v2', 'report', 'img')
OUT = os.path.join(ROOT, 'docs', 'v2', 'report')
CACHE = os.path.join(os.environ.get('TEMP', 'C:/gt'), 'guide_v2_cache')
os.makedirs(CACHE, exist_ok=True)
NAVY, LIME, CREAM, INK = '#061A31', '#BDF642', '#FCF7EB', '#1d2733'
FB = 'C:/Windows/Fonts/NanumGothicBold.ttf'
FR = 'C:/Windows/Fonts/NanumGothic.ttf'
BASENAME = '노포패스포트_버전2_만들기'


def find(name):
    for d in (SCREENS, IMG):
        p = os.path.join(d, name)
        if os.path.exists(p):
            return p
    raise FileNotFoundError(name)


# ------------------------------------------------------------------ 다이어그램
def _node(i, x, y, w, h, t, sub, fill, fg='#061A31'):
    return dict(id=i, x=x, y=y, w=w, h=h, t=t, sub=sub, fill=fill, fg=fg)


def diagrams():
    D = {}
    names = [('기획', '서비스 정의\n규칙은 숫자로\n화면 목록', '#E8F3C4'), ('디자인', 'Stitch 51화면\n사용자 시안 34장\n앱 아이콘·브랜드', '#F7E7B6'),
             ('데이터 수집', '백년가게·오래가게\n1,277곳\n좌표·메뉴 사진', '#D9ECF7'), ('개발', '스토리 하나씩\n테스트 먼저\n실기기 설치', '#F4D7D1'),
             ('테스트', '단위 테스트 48파일\n자동 탭 점검\nmonkey 3,000회', '#E1D8F2'), ('Firebase 연결', '로그인 · Firestore\nAI(Gemini)\n구글 로그인 설정', '#FFE3C2'),
             ('세부 수정', '사용자 점검\n프롬프트 → 결과\n다시 확인', '#CFEFE3')]
    W, G = 420, 50
    nodes = []
    for i in range(4):
        nodes.append(_node(i, 20 + i * (W + G), 40, W, 230, '%d. %s' % (i + 1, names[i][0]), names[i][1], names[i][2]))
    for j, i in enumerate(range(4, 7)):
        col = 3 - j
        nodes.append(_node(i, 20 + col * (W + G), 340, W, 230, '%d. %s' % (i + 1, names[i][0]), names[i][1], names[i][2]))
    nodes.append(_node(9, 20, 640, 1810, 140, '※ 4~7단계는 한 번으로 끝나지 않았다', '사용자가 써 보고 고칠 것을 말하면 개발 → 테스트 → 확인을 되풀이했다 (세부 수정 내역 참고)', '#FBE3DF'))
    D['process'] = dict(w=1850, h=810, title=None, f1=36, f2=27, step=38, nodes=nodes,
                        edges=[(0, 1), (1, 2), (2, 3), (3, 4), (4, 5), (5, 6)], poly=[])
    # 구조
    L = [('화면 (Jetpack Compose)', 'feature/ home · explore · map · passport · verify · ai · my · settings · notifications', '#F7E7B6'),
         ('ViewModel', '상태(StateFlow) ← 화면의 행동. 단방향 흐름, SavedStateHandle로 단계·시도 횟수 보존', '#E8F3C4'),
         ('domain 규칙 (테스트를 먼저)', 'TitlePolicy · VerificationConfig · DailyLimitPolicy · RoulettePicker · TasteRecommender', '#D9ECF7')]
    nodes = [_node(i, 30, 30 + i * 150, 1330, 120, t, s, c) for i, (t, s, c) in enumerate(L)]
    dx = [('Room', '카탈로그 1,277곳\n내 방문 기록', '#F4D7D1'), ('DataStore', '설정 · 취향\n저장한 노포', '#F4D7D1'),
          ('Cloud Firestore', '프로필 · 닉네임\n후기 · 방문(서버)', '#FFE3C2'), ('Firebase AI', 'Gemini: 룰렛 · 도우미\n번역 · 촬영 안내', '#FFE3C2')]
    for i, (t, s, c) in enumerate(dx):
        nodes.append(_node(10 + i, 30 + i * 335, 510, 310, 150, t, s, c))
    pf = [('CameraX', '영수증·간판 촬영', '#E1D8F2'), ('ML Kit 한국어 OCR', '글자 읽기(앱에 포함)', '#E1D8F2'),
          ('Fused Location', '가게 30m 확인', '#E1D8F2'), ('네이버 지도 SDK', '지도 · 마커', '#E1D8F2')]
    for i, (t, s, c) in enumerate(pf):
        nodes.append(_node(20 + i, 30 + i * 335, 720, 310, 120, t, s, c))
    nodes.append(_node(30, 1420, 30, 330, 810, 'AppContainer', '수동 주입\n(Hilt 없이)\n\n카탈로그 · 방문 · 설정\n인증 · AI · 번역을\n한 곳에서 만들어\n화면에 넘긴다', '#CFEFE3'))
    D['arch'] = dict(w=1780, h=870, title=None, f1=32, f2=25, step=34, nodes=nodes,
                     edges=[(0, 1), (1, 2), (2, 10), (2, 11), (2, 12), (2, 13)], poly=[])
    # 데이터
    W, G = 400, 70
    X = lambda c: 20 + c * (W + G)
    n = [_node(0, X(0), 40, W, 230, '공식 자료', '소상공인마당 백년가게\n서울 오래가게 발표\n공공데이터포털 목록', '#F7E7B6'),
         _node(1, X(1), 40, W, 230, '수집 원본', 'data/raw/*.jsonl\n가게당 한 줄', '#E8F3C4'),
         _node(2, X(2), 40, W, 230, 'build_catalog.py', '이름 별칭 · 업력\n음식점 여부 정리', '#D9ECF7'),
         _node(6, X(3), 40, W, 100, 'geocode.py', '주소 → 좌표', '#E1D8F2'),
         _node(7, X(3), 170, W, 100, '메뉴 사진 받기', 'fetch_menu_photos.py (CC)', '#E1D8F2'),
         _node(3, X(2), 340, W, 230, 'stores.json', '앱용 1,277곳\n+ csv · db · stats', '#F4D7D1'),
         _node(4, X(1), 340, W, 230, '앱 첫 실행', 'Room에 적재\n(버전이 바뀌면 다시)', '#FFE3C2'),
         _node(5, X(0), 340, W, 230, '화면', '지도 · 탐색 · 룰렛\n방문 인증 판정', '#CFEFE3')]
    D['data'] = dict(w=1850, h=600, title=None, f1=34, f2=26, step=36, nodes=n,
                     edges=[(0, 1), (1, 2), (2, 3), (3, 4), (4, 5), (6, 2), (7, 2)], poly=[])
    return D


def _edge_pts(a, b):
    ax, ay, aw, ah = a['x'], a['y'], a['w'], a['h']
    bx, by, bw, bh = b['x'], b['y'], b['w'], b['h']
    acx, acy, bcx, bcy = ax + aw / 2, ay + ah / 2, bx + bw / 2, by + bh / 2
    same_row = abs(acy - bcy) < max(ah, bh) / 2 + 5 or (by < ay + ah and ay < by + bh)
    if same_row and (bx >= ax + aw - 5 or bx + bw <= ax + 5):
        y = min(max(acy, by + 20), by + bh - 20) if (ah != bh) else acy
        y = bcy if bh < ah else acy
        if bx > ax:
            return (ax + aw, y), (bx, y)
        return (ax, y), (bx + bw, y)
    x = acx if (bx <= acx <= bx + bw) else bcx
    if by > ay:
        return (x, ay + ah), (x, by)
    return (x, ay), (x, by + bh)


def _wrap(draw_or_none, text, font, maxw):
    out = []
    for para in text.split('\n'):
        line = ''
        for ch in para:
            if font.getlength(line + ch) > maxw and line:
                out.append(line)
                line = ch
            else:
                line += ch
        out.append(line)
    return out


def _fonts(spec):
    return ImageFont.truetype(FB, spec['f1']), ImageFont.truetype(FR, spec['f2'])


def render_png(spec, path):
    W, Hh = spec['w'], spec['h']
    im = Image.new('RGB', (W, Hh), '#FFFFFF')
    d = ImageDraw.Draw(im)
    f1, f2 = _fonts(spec)
    byid = {n['id']: n for n in spec['nodes']}
    for a, b in spec['edges']:
        p, q = _edge_pts(byid[a], byid[b])
        d.line([p, q], fill=NAVY, width=6)
        dx, dy = q[0] - p[0], q[1] - p[1]
        L = max((dx ** 2 + dy ** 2) ** .5, 1)
        ux, uy = dx / L, dy / L
        d.polygon([q, (q[0] - ux * 22 - uy * 11, q[1] - uy * 22 + ux * 11), (q[0] - ux * 22 + uy * 11, q[1] - uy * 22 - ux * 11)], fill=NAVY)
    for n in spec['nodes']:
        d.rounded_rectangle((n['x'], n['y'], n['x'] + n['w'], n['y'] + n['h']), 24, fill=n['fill'], outline=NAVY, width=3)
        tw = d.textlength(n['t'], font=f1)
        d.text((n['x'] + (n['w'] - tw) / 2, n['y'] + 14), n['t'], font=f1, fill=n['fg'])
        yy = n['y'] + 14 + spec['f1'] + 22
        for line in _wrap(d, n['sub'], f2, n['w'] - 30):
            lw = d.textlength(line, font=f2)
            d.text((n['x'] + (n['w'] - lw) / 2, yy), line, font=f2, fill='#33414F')
            yy += spec['step']
    im.save(path)


def render_svg(spec):
    W, Hh = spec['w'], spec['h']
    f1, f2 = _fonts(spec)
    o = ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 %d %d" role="img" style="width:100%%;height:auto;font-family:\'Noto Sans KR\',\'Malgun Gothic\',sans-serif">' % (W, Hh),
         '<defs><marker id="ah" markerWidth="12" markerHeight="12" refX="9" refY="5" orient="auto"><path d="M0,0 L10,5 L0,10 z" fill="#061A31"/></marker></defs>']
    byid = {n['id']: n for n in spec['nodes']}
    for a, b in spec['edges']:
        p, q = _edge_pts(byid[a], byid[b])
        o.append('<line x1="%d" y1="%d" x2="%d" y2="%d" stroke="#061A31" stroke-width="6" marker-end="url(#ah)"/>' % (p[0], p[1], q[0], q[1]))
    for n in spec['nodes']:
        o.append('<g class="nd"><rect x="%d" y="%d" width="%d" height="%d" rx="24" fill="%s" stroke="#061A31" stroke-width="3"/>' % (n['x'], n['y'], n['w'], n['h'], n['fill']))
        o.append('<text x="%d" y="%d" font-size="%d" font-weight="800" fill="%s" text-anchor="middle">%s</text>' % (n['x'] + n['w'] / 2, n['y'] + 14 + spec['f1'] * 0.85, spec['f1'], n['fg'], H.escape(n['t'])))
        yy = n['y'] + 14 + spec['f1'] + 22 + spec['f2'] * 0.8
        for line in _wrap(None, n['sub'], f2, n['w'] - 30):
            o.append('<text x="%d" y="%d" font-size="%d" fill="#33414F" text-anchor="middle">%s</text>' % (n['x'] + n['w'] / 2, yy, spec['f2'], H.escape(line)))
            yy += spec['step']
        o.append('</g>')
    o.append('</svg>')
    return ''.join(o)


DIAGS = diagrams()
for k, v in DIAGS.items():
    render_png(v, os.path.join(CACHE, 'diagram_%s.png' % k))
    render_png(v, os.path.join(IMG, 'guide_diagram_%s.png' % k))


# ------------------------------------------------------------------ 이미지 준비
def jpg_path(name, maxw, q=84):
    p = find(name)
    key = os.path.join(CACHE, '%d_%s.jpg' % (maxw, os.path.splitext(os.path.basename(p))[0]))
    if not os.path.exists(key) or os.path.getmtime(key) < os.path.getmtime(p):
        im = Image.open(p).convert('RGB')
        if im.width > maxw:
            im = im.resize((maxw, int(im.height * maxw / im.width)), Image.LANCZOS)
        im.save(key, quality=q, optimize=True)
    return key


def data_uri(path):
    ext = os.path.splitext(path)[1].lower()
    mime = 'image/png' if ext == '.png' else 'image/jpeg'
    return 'data:%s;base64,%s' % (mime, base64.b64encode(open(path, 'rb').read()).decode())


TITLES = [('첫술 식객', 0), ('골목 단골', 100), ('국밥 순례자', 200), ('시장통 어르신', 300), ('가마솥 감별사', 400), ('노포 길잡이', 500),
          ('맛의 장인', 600), ('세월 지킴이', 700), ('미식 도인', 800), ('전설의 식객', 900), ('천년 미식가', 1000)]


def commit_count():
    try:
        return subprocess.run(['git', '-C', ROOT, 'rev-list', '--count', 'HEAD'], capture_output=True, text=True).stdout.strip() or '189'
    except Exception:
        return '189'


STATS = [('1,277', '앱에 담은 노포'), ('11', '칭호 단계 (세월 100년마다)'), ('50', '실기기 화면 캡처'), ('48', '단위 테스트 파일'),
         ('3,000', 'monkey 임의 조작 (앱 종료 0)'), (commit_count(), '커밋')]


# ------------------------------------------------------------------ DOCX
def build_docx(path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    FONT = '맑은 고딕'

    def setfont(style, size=None, bold=None, color=None, name=FONT):
        style.font.name = name
        rpr = style.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.append(rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
            rf.set(qn(a), name)
        if size:
            style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if color:
            style.font.color.rgb = RGBColor.from_string(color)

    setfont(doc.styles['Normal'], 10.5)
    doc.styles['Normal'].paragraph_format.space_after = Pt(4)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.25
    setfont(doc.styles['Title'], 26, True, '061A31')
    setfont(doc.styles['Heading 1'], 17, True, '061A31')
    setfont(doc.styles['Heading 2'], 13, True, '2A4A12')
    for nm in ('Heading 1', 'Heading 2'):
        doc.styles[nm].paragraph_format.keep_with_next = True
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(20)
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(12)
    setfont(doc.styles['List Bullet'], 10.5)
    setfont(doc.styles['List Number'], 10.5)

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill)
        tcPr.append(shd)

    def borders(table, color='C8CBB8', sz=4):
        tbl = table._tbl
        pr = tbl.tblPr
        b = OxmlElement('w:tblBorders')
        for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            x = OxmlElement('w:' + e)
            x.set(qn('w:val'), 'single')
            x.set(qn('w:sz'), str(sz))
            x.set(qn('w:space'), '0')
            x.set(qn('w:color'), color)
            b.append(x)
        pr.append(b)

    def no_split(row):
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement('w:cantSplit')
        trPr.append(cs)

    def para(text, style=None, size=None, bold=False, color=None, align=None, cell=None, italic=False, mono=False):
        p = (cell.add_paragraph() if cell is not None else doc.add_paragraph(style=style))
        r = p.add_run(text)
        if size:
            r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if mono:
            r.font.name = 'Consolas'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        if align:
            p.alignment = align
        return p

    def box(lines, fill, mono=False, size=10, italic=False):
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(t, 'D4D8C0')
        c = t.rows[0].cells[0]
        shade(c, fill)
        c.paragraphs[0].text = ''
        first = True
        for ln in lines:
            p = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            r = p.add_run(ln)
            r.font.size = Pt(size)
            r.italic = italic
            if mono:
                r.font.name = 'Consolas'
                r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
            p.paragraph_format.space_after = Pt(1)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # 바닥글 쪽 번호
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('노포 패스포트 만들기 (버전2)  ·  ')
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string('7A7F70')
    r2 = fp.add_run()
    r2.font.size = Pt(8.5)
    for t, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if t:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), t)
        else:
            e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve')
            e.text = txt
        r2._r.append(e)

    figno = [0]

    def caption(text, cell=None, align=WD_ALIGN_PARAGRAPH.CENTER):
        return para(text, size=8.5, color='6B7060', align=align, cell=cell)

    first_title = True
    for blk in content.B:
        k = blk[0]
        if k == 'title':
            icon = doc.add_paragraph()
            icon.add_run().add_picture(os.path.join(ROOT, 'design', 'brand', 'play-store-icon-512.png'), width=Cm(2.8))
            doc.add_paragraph(blk[1], style='Title')
        elif k == 'title2':
            doc.add_page_break() if not first_title and False else None
            doc.add_paragraph(blk[1], style='Title')
        elif k == 'lines':
            for ln in blk[1]:
                para(ln, size=10)
        elif k == 'stats':
            t = doc.add_table(rows=2, cols=len(STATS))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            borders(t, 'D4D8C0')
            for i, (n, lab) in enumerate(STATS):
                c0, c1 = t.rows[0].cells[i], t.rows[1].cells[i]
                shade(c0, 'E8F3C4')
                shade(c1, 'F7FAEC')
                c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = c0.paragraphs[0].add_run(n)
                rr.bold = True
                rr.font.size = Pt(15)
                c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                r3 = c1.paragraphs[0].add_run(lab)
                r3.font.size = Pt(8)
            doc.add_paragraph()
        elif k == 'h1':
            doc.add_paragraph(blk[1], style='Heading 1')
        elif k == 'h2':
            doc.add_paragraph(blk[1], style='Heading 2')
        elif k == 'p':
            para(blk[1])
        elif k == 'ul':
            for t in blk[1]:
                doc.add_paragraph(t, style='List Bullet')
        elif k == 'ol':
            for i, t in enumerate(blk[1], 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.add_run('%d. ' % i).bold = True
                p.add_run(t)
        elif k == 'code':
            box(blk[1].split('\n'), 'F1F1EA', mono=True, size=8.5)
        elif k == 'note':
            box([blk[1]], 'FFF4D6', size=9.5)
        elif k == 'prompt':
            box(blk[1], 'EAF4C4', size=10)
        elif k == 'table':
            hdr, rows = blk[1], blk[2]
            t = doc.add_table(rows=1, cols=len(hdr))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            borders(t)
            for i, h in enumerate(hdr):
                c = t.rows[0].cells[i]
                shade(c, '061A31')
                c.paragraphs[0].text = ''
                rr = c.paragraphs[0].add_run(h)
                rr.bold = True
                rr.font.size = Pt(9.5)
                rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            for ri, row in enumerate(rows):
                cells = t.add_row().cells
                no_split(t.rows[-1])
                for i, v in enumerate(row):
                    cells[i].paragraphs[0].text = ''
                    rr = cells[i].paragraphs[0].add_run(v)
                    rr.font.size = Pt(9.5)
                    if ri % 2:
                        shade(cells[i], 'F6F8EE')
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif k == 'titles':
            t = doc.add_table(rows=2, cols=len(TITLES))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            borders(t)
            shades = ['EFF6D5', 'E8F3C4', 'DFEEAA', 'D6E88F', 'CDE274', 'F7E7B6', 'F4D7A0', 'F0C98A', 'EDB87A', 'E8A86E', 'F2D25A']
            for i, (nm, y) in enumerate(TITLES):
                a, b = t.rows[0].cells[i], t.rows[1].cells[i]
                shade(a, shades[i])
                shade(b, shades[i])
                a.paragraphs[0].alignment = b.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                ra = a.paragraphs[0].add_run(nm)
                ra.bold = True
                ra.font.size = Pt(7.5)
                rb = b.paragraphs[0].add_run('%d년' % y)
                rb.font.size = Pt(8)
            caption('그림 · 칭호 11단계 (누적 세월 기준)')
        elif k == 'shots':
            items = blk[1]
            n = len(items)
            wcm = min(4.1, 16.6 / n)
            t = doc.add_table(rows=2, cols=n)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            no_split(t.rows[0])
            for i, (fn, cap) in enumerate(items):
                c = t.rows[0].cells[i]
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                c.paragraphs[0].add_run().add_picture(jpg_path(fn, 540), width=Cm(wcm))
                cc = t.rows[1].cells[i]
                cc.paragraphs[0].text = ''
                cc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = cc.paragraphs[0].add_run(cap)
                rr.font.size = Pt(8)
                rr.font.color.rgb = RGBColor.from_string('6B7060')
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif k == 'pair':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(jpg_path(blk[1], 1100), width=Cm(11.5))
            caption(blk[2])
        elif k == 'diagram':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(os.path.join(CACHE, 'diagram_%s.png' % blk[1]), width=Cm(16.6))
            caption(blk[2])
        elif k == 'pagebreak':
            doc.add_page_break()
    doc.core_properties.title = '노포 패스포트 만들기 (버전2)'
    doc.core_properties.author = '노포 패스포트'
    doc.save(path)


# ------------------------------------------------------------------ HTML
CSS = """
:root{--navy:#061A31;--lime:#BDF642;--cream:#FCF7EB;--ink:#1d2733;--mut:#5d6675;--line:#e3ddc9;--card:#fff}
*{box-sizing:border-box}html{scroll-behavior:smooth}
body{margin:0;font-family:'Noto Sans KR','Malgun Gothic','Apple SD Gothic Neo',sans-serif;color:var(--ink);background:var(--cream);line-height:1.7;word-break:keep-all}
.layout{display:flex;max-width:1280px;margin:0 auto}
nav.toc{position:sticky;top:0;align-self:flex-start;height:100vh;overflow:auto;width:250px;flex:none;padding:26px 14px;border-right:1px solid var(--line);background:#fffdf6}
nav.toc .brand{display:flex;gap:10px;align-items:center;margin-bottom:16px}nav.toc img{width:44px;height:44px;border-radius:12px}
nav.toc b{font-size:15px;line-height:1.25}
nav.toc a{display:block;padding:6px 10px;border-radius:10px;color:var(--ink);text-decoration:none;font-size:14px}
nav.toc a:hover,nav.toc a.on{background:var(--lime);font-weight:700}
nav.toc .sub{display:flex;flex-wrap:wrap;gap:4px;padding:2px 8px 8px}
nav.toc .sub a{padding:2px 8px;font-size:12px;background:#f1ecd9}
main{flex:1;min-width:0;padding:0 34px 80px}
.hero{background:linear-gradient(135deg,#061A31,#0b2a4a 62%,#12395f);color:#fff;border-radius:0 0 28px 28px;padding:44px 40px;margin:0 -34px 28px;display:flex;gap:30px;align-items:center;flex-wrap:wrap}
.hero img{width:130px;height:130px;border-radius:30px;box-shadow:0 10px 30px #0008}
.hero h1{font-size:40px;margin:0 0 6px;letter-spacing:-1px;color:#fff;display:block}.hero h1::before{display:none}.hero h1 .acc{color:var(--lime)}.hero p{margin:4px 0;color:#d6e2ef}
.pill{display:inline-block;background:var(--lime);color:var(--navy);font-weight:800;border-radius:999px;padding:2px 14px;font-size:14px;margin-left:8px;vertical-align:middle}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin:18px 0}
.stat{background:var(--card);border:1px solid var(--line);border-radius:18px;padding:16px;text-align:center}
.stat b{display:block;font-size:30px;color:var(--navy);letter-spacing:-1px}.stat span{font-size:13px;color:var(--mut)}
.tools{background:var(--card);border:1px solid var(--line);border-radius:18px;padding:18px 22px;margin:14px 0}
.tools p{margin:6px 0;font-size:14.5px}
h1{font-size:30px;margin:54px 0 8px;color:var(--navy);display:flex;align-items:center;gap:12px;scroll-margin-top:14px}
h1::before{content:"";width:12px;height:34px;border-radius:6px;background:var(--lime);flex:none}
h1.entry{font-size:22px;margin-top:40px;background:#fff;border:1px solid var(--line);border-radius:14px;padding:8px 16px}
h2{font-size:20px;margin:30px 0 6px;color:#2A4A12}
h2.lab{display:inline-block;font-size:13px;letter-spacing:1px;color:#fff;background:#3a6df0;border-radius:999px;padding:1px 14px;margin:14px 0 6px}
h2.lab.res{background:#2f9e4f}
p{margin:8px 0}ul,ol{padding-left:22px}li{margin:3px 0}
.prompt{background:#eaf4c4;border-left:6px solid #8fb92b;border-radius:12px;padding:12px 18px;font-size:15px}
.prompt div{margin:2px 0}
.note{background:#fff4d6;border-left:6px solid #e6b422;border-radius:12px;padding:10px 16px;font-size:14.5px;margin:14px 0}
pre.code{background:#101c2b;color:#dff6a8;border-radius:12px;padding:14px 18px;font:13px/1.6 Consolas,'D2Coding',monospace;overflow:auto;white-space:pre-wrap}
table{border-collapse:separate;border-spacing:0;width:100%;margin:12px 0;background:#fff;border:1px solid var(--line);border-radius:14px;overflow:hidden;font-size:14px}
th{background:var(--navy);color:#fff;text-align:left;padding:9px 12px;font-size:13.5px}
td{padding:8px 12px;border-top:1px solid var(--line);vertical-align:top}tr:nth-child(even) td{background:#fafcf1}
.shots{display:flex;flex-wrap:wrap;gap:16px;margin:14px 0}
.shots figure{margin:0;width:210px;text-align:center}
.shots img{width:100%;border-radius:22px;border:6px solid #14202e;background:#14202e;cursor:zoom-in;box-shadow:0 8px 22px #0002}
figcaption,.cap{font-size:12.5px;color:var(--mut);margin-top:6px;text-align:center}
.pair{margin:16px 0;text-align:center}.pair img{max-width:100%;width:760px;border-radius:16px;border:1px solid var(--line);cursor:zoom-in}
.pair .lbl{display:flex;justify-content:center;gap:230px;font-size:13px;font-weight:700;color:var(--navy);margin-bottom:6px}
.diag{background:#fff;border:1px solid var(--line);border-radius:18px;padding:10px;margin:14px 0}
.diag svg .nd{transition:transform .2s}.diag svg .nd:hover{transform:translateY(-3px)}
.ladder{display:flex;align-items:flex-end;gap:6px;margin:16px 0 4px;min-height:170px}
.ladder div{flex:1;border-radius:10px 10px 4px 4px;padding:6px 2px;text-align:center;font-size:12px;font-weight:700;color:#061A31;border:2px solid #061A31}
.ladder small{display:block;font-weight:400;font-size:11px}
.grade{display:flex;gap:18px;margin:6px 0 14px;flex-wrap:wrap}.grade span{display:flex;align-items:center;gap:8px;font-size:14px}
.grade i{width:26px;height:26px;border-radius:50%;display:inline-block;border:3px solid}
dialog{border:none;background:transparent;padding:0;max-width:92vw;max-height:92vh}dialog::backdrop{background:#000b}
dialog img{max-width:92vw;max-height:90vh;border-radius:16px}
@media(max-width:900px){nav.toc{display:none}main{padding:0 16px 60px}.hero{margin:0 -16px 24px}.shots figure{width:44%}}
@media print{nav.toc{display:none}.hero{border-radius:0}h1{break-after:avoid}.shots figure,table,.pair{break-inside:avoid}}
"""

JS = """
const links=[...document.querySelectorAll('nav.toc a[href^="#"]')];
const secs=links.map(a=>document.querySelector(a.getAttribute('href'))).filter(Boolean);
addEventListener('scroll',()=>{let cur=0;secs.forEach((s,i)=>{if(s.getBoundingClientRect().top<140)cur=i});links.forEach(a=>a.classList.remove('on'));if(links[cur])links[cur].classList.add('on')});
const dlg=document.getElementById('zoom'),zi=dlg.querySelector('img');
document.querySelectorAll('.shots img,.pair img').forEach(i=>i.addEventListener('click',()=>{zi.src=i.src;dlg.showModal()}));
dlg.addEventListener('click',()=>dlg.close());
"""


def build_html(path):
    esc = H.escape
    out = []
    nav = []
    ids = {}
    cur_sub = None
    sid = [0]
    sections = []

    # 목차 구성
    for blk in content.B:
        if blk[0] == 'h1' and not blk[1].startswith('# '):
            sid[0] += 1
            sections.append((blk[1], 'sec%d' % sid[0], []))
        elif blk[0] == 'title2':
            sid[0] += 1
            sections.append((blk[1], 'sec%d' % sid[0], []))
        elif blk[0] == 'h1' and blk[1].startswith('# ') and sections:
            sections[-1][2].append(blk[1])
    icon_uri = data_uri(jpg_path_icon())
    nav.append('<nav class="toc"><div class="brand"><img src="%s" alt=""><b>노포 패스포트<br>만들기 (버전2)</b></div>' % icon_uri)
    for t, i, subs in sections:
        nav.append('<a href="#%s">%s</a>' % (i, esc(t)))
        if subs:
            nav.append('<div class="sub">%s</div>' % ''.join('<a href="#e%s">%s</a>' % (esc(s[2:]).replace('.', '-'), esc(s)) for s in subs))
    nav.append('</nav>')

    out.append('<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
               '<title>노포 패스포트 만들기 (버전2)</title><style>%s</style></head><body><div class="layout">' % CSS)
    out.extend(nav)
    out.append('<main>')
    sec_iter = iter(sections)
    pending_title = None
    for blk in content.B:
        k = blk[0]
        if k == 'title':
            out.append('<header class="hero"><img src="%s" alt="앱 아이콘"><div><h1>노포 패스포트 <span class="acc">만들기</span><span class="pill">버전2</span></h1>'
                       '<p>기획 · 디자인 · 데이터 · 개발 · 테스트 · Firebase · 세부 수정까지, 실제로 한 순서 그대로</p></div></header>' % icon_uri)
        elif k == 'lines':
            L = blk[1]
            out.append('<div class="tools"><p><b>%s</b></p>%s</div>' % (esc(L[0]), ''.join('<p>%s</p>' % linkify(esc(x)) for x in L[1:])))
        elif k == 'stats':
            out.append('<div class="stats">%s</div>' % ''.join('<div class="stat"><b>%s</b><span>%s</span></div>' % (n, esc(l)) for n, l in STATS))
        elif k == 'h1':
            if blk[1].startswith('# '):
                eid = 'e' + blk[1][2:].replace('.', '-')
                out.append('<h1 class="entry" id="%s">%s</h1>' % (eid, esc(blk[1])))
            else:
                t, i, _ = next(sec_iter)
                out.append('<h1 id="%s">%s</h1>' % (i, esc(blk[1])))
        elif k == 'title2':
            t, i, _ = next(sec_iter)
            out.append('<h1 id="%s">%s</h1>' % (i, esc(blk[1])))
        elif k == 'h2':
            if blk[1] in ('PROMPT', 'RESULT'):
                out.append('<h2 class="lab%s">%s</h2>' % (' res' if blk[1] == 'RESULT' else '', blk[1]))
            else:
                out.append('<h2>%s</h2>' % esc(blk[1]))
        elif k == 'p':
            out.append('<p>%s</p>' % esc(blk[1]))
        elif k == 'ul':
            out.append('<ul>%s</ul>' % ''.join('<li>%s</li>' % esc(t) for t in blk[1]))
        elif k == 'ol':
            out.append('<ol>%s</ol>' % ''.join('<li>%s</li>' % esc(t) for t in blk[1]))
        elif k == 'code':
            out.append('<pre class="code">%s</pre>' % esc(blk[1]))
        elif k == 'note':
            out.append('<div class="note">%s</div>' % esc(blk[1]))
        elif k == 'prompt':
            out.append('<div class="prompt">%s</div>' % ''.join('<div>%s</div>' % esc(x) for x in blk[1]))
        elif k == 'table':
            out.append('<table><thead><tr>%s</tr></thead><tbody>%s</tbody></table>' % (
                ''.join('<th>%s</th>' % esc(h) for h in blk[1]),
                ''.join('<tr>%s</tr>' % ''.join('<td>%s</td>' % esc(c) for c in r) for r in blk[2])))
        elif k == 'titles':
            cols = ['#EFF6D5', '#E8F3C4', '#DFEEAA', '#D6E88F', '#CDE274', '#F7E7B6', '#F4D7A0', '#F0C98A', '#EDB87A', '#E8A86E', '#F2D25A']
            out.append('<div class="ladder">%s</div><div class="cap">칭호 11단계 · 세월 100년마다 한 단계, 1,000년이면 천년 미식가</div>' % ''.join(
                '<div style="height:%dpx;background:%s">%s<small>%d년</small></div>' % (60 + i * 11, cols[i], esc(n), y) for i, (n, y) in enumerate(TITLES)))
            out.append('<div class="grade"><span><i style="border-color:#2E8B8B;background:#bfe3e3"></i>30~49년 청록 도장</span><span><i style="border-color:#B4332A;background:#f1c9c4"></i>50~99년 인주 빨강</span><span><i style="border-color:#C9A24A;background:#f6e5a8"></i>100년 이상 금빛</span></div>')
        elif k == 'shots':
            out.append('<div class="shots">%s</div>' % ''.join(
                '<figure><img loading="lazy" src="%s" alt="%s"><figcaption>%s</figcaption></figure>' % (data_uri(jpg_path(fn, 400, 80)), esc(cap), esc(cap)) for fn, cap in blk[1]))
        elif k == 'pair':
            out.append('<div class="pair"><div class="lbl"><span>시안</span><span>앱 (실기기)</span></div><img loading="lazy" src="%s" alt="%s"><div class="cap">%s</div></div>' % (
                data_uri(jpg_path(blk[1], 1000, 80)), esc(blk[2]), esc(blk[2])))
        elif k == 'diagram':
            out.append('<div class="diag">%s<div class="cap">%s</div></div>' % (render_svg(DIAGS[blk[1]]), esc(blk[2])))
        elif k == 'pagebreak':
            pass
    out.append('</main></div><dialog id="zoom"><img alt=""></dialog><script>%s</script></body></html>' % JS)
    open(path, 'w', encoding='utf-8').write('\n'.join(out))


def jpg_path_icon():
    p = os.path.join(ROOT, 'design', 'brand', 'play-store-icon-512.png')
    key = os.path.join(CACHE, 'icon_160.jpg')
    Image.open(p).convert('RGB').resize((160, 160), Image.LANCZOS).save(key, quality=90)
    return key


def linkify(s):
    import re
    return re.sub(r'(https?://[^\s)]+)', r'<a href="\1" style="color:#bdf642">\1</a>', s) if False else re.sub(r'(https?://[^\s)]+)', r'<a href="\1" target="_blank" rel="noopener" style="color:#0b6b2a">\1</a>', s)


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    build_docx(os.path.join(OUT, BASENAME + '.docx'))
    build_html(os.path.join(OUT, BASENAME + '.html'))
    for ext in ('docx', 'html'):
        p = os.path.join(OUT, BASENAME + '.' + ext)
        print(p, os.path.getsize(p) // 1024, 'KB')
